#!/usr/bin/env python3
"""
Funnel delt i A og B, en fane til hver, plus de fulde lister.

  Funnel A         ejer bor på adressen: fra ejerforeninger til køb, med handling og link til brev og mails
  Funnel B         ejer bor et andet sted: samme tragt for sig selv
  Ejerforeninger   alle foreninger, hvor de kommer fra, og hvorfor de er med eller ikke med
  Ejere            alle lejligheder på brevlisten (segment A og B), ejer, kontakt, breve og mails sendt
  Breve            hvem der har fået brev, hvilket budskab og hvor mange gange

Layoutet følger Jacobs egen redigering af v6 (uden søjle-tragten og uden undertitlen).
Indeholder personoplysninger, så filen ligger kun i vaulten.

Kør:  python3 scripts/funnel/bygfunnel_ab.py --ejere ejernavne.json --kontakt leadkontakt.txt --mails mailhistorik.txt
"""
import argparse
import sys
from pathlib import Path

from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).parent))
from bygfunnel import VAULT  # noqa: E402
from bygfunnel_simpel import BREV, BUDSKAB, LINK_FONT, LINKS, byg  # noqa: E402
from bygliste import byg_lister  # noqa: E402
from mstyle import PCT, fml, hard, header_row, lbl, note, title, widths  # noqa: E402

HVEM = {'A': 'ejer bor på adressen', 'B': 'ejer bor et andet sted'}
BREVRAEKKE = {'A': 5, 'B': 6}   # rækken i Breve-fanen

# (trin, nøgle i tællingen, handling, [(dokument, hvornår)])
TRIN = [
    ('Lead', 'lead', 'Send booking-mailen senest dagen efter, at de har brugt beregneren.',
     [('Mail 01 Booking', 'Senest dagen efter')]),
    ('Booking sendt', 'booking', 'Følg op i samme tråd. Uden svar efter tredje opfølgning parkeres leadet.',
     [('Mail 02 Opfølgning 1 dag 3', 'Dag 3'), ('Mail 03 Opfølgning 2 dag 8', 'Dag 8'), ('Mail 04 Opfølgning 3 dag 15', 'Dag 15')]),
    ('Besigtigelse aftalt', 'aftalt', 'Bekræft tiden skriftligt, og påmind dagen før. Kør derud med scriptet.',
     [('Mail 05 Bekræftelse', 'Straks'), ('Mail 06 Påmindelse', 'Dagen før'), ('Script Besigtigelse', 'På dagen')]),
    ('Besigtigelse afholdt', 'afholdt', 'Giv et skriftligt bud senest dagen efter.',
     [('Mail 07 Tak for i dag', 'Samme dag')]),
    ('Bud afgivet', 'bud', 'Send kontantbud, som gælder i 30 dage, og følg op. Tæller også dem, der siden blev uenige om prisen.',
     [('Mail 08 Bud', 'Dagen efter besigtigelsen'), ('Mail 09 Bud-opfølgning dag 3', 'Dag 3')]),
    ('Købt', 'koebt', 'Send de næste skridt, når I er blevet enige.',
     [('Mail 14 Sådan foregår handlen', 'Når I er enige')]),
]


def funnel_fane(wb, seg, ctx):
    ws = wb.create_sheet(f'Funnel {seg}')
    tal = ctx['tal'][seg]
    br = BREVRAEKKE[seg]
    title(ws, f'Funnel {seg}: {HVEM[seg]}, fra ejerforeninger til køb', None, 7)
    widths(ws, {'A': 34, 'B': 11, 'C': 12, 'D': 10, 'E': 56, 'F': 32, 'G': 24})

    header_row(ws, 4, ['Fra ejerforeninger til brev', 'Antal', '% af forrige', '', 'Hvad tallet er'])
    ws.row_dimensions[4].height = 26
    top = [
        ('Ejerforeninger identificeret', ctx['antal_foreninger'], 'Alle foreninger i registret på /foreninger'),
        ('heraf udvalgt', ctx['udvalgt'], 'Status Målgruppe: dem vi vil købe i'),
        ('heraf størrelse 20–80 kvm', ctx['stoerrelse'], 'Boliger i de udvalgte, målt i BBR'),
        ('kan få brev', '=Breve!C7', 'Ejeren er kendt i Resights, og koncernens egne lejligheder er trukket fra'),
        (f'Segment {seg}: {HVEM[seg]}', f'=Breve!C{br}',
         'Brevet når ejeren selv' if seg == 'A' else 'Brevet lander hos lejeren. Ejerens egen adresse står i fanen Ejere'),
    ]
    for i, (navn, v, kom) in enumerate(top):
        r = 5 + i
        lbl(ws, f'A{r}', navn, bold=i in (0, 3), indent=1 if i in (1, 2) else 2 if i == 4 else 0)
        (fml if isinstance(v, str) else hard)(ws, f'B{r}', v)
        if i in (1, 3):
            fml(ws, f'C{r}', f'=IF(B{r-1}>0,B{r}/B{r-1},"–")', fmt=PCT)
        elif i == 4:
            fml(ws, f'C{r}', f'=IF($B$8>0,B{r}/$B$8,"–")', fmt=PCT)
        note(ws, f'E{r}', kom)
        for c in 'BC':
            ws[f'{c}{r}'].alignment = Alignment(horizontal='right', vertical='center')
        ws.row_dimensions[r].height = 22

    hr = 12
    header_row(ws, hr, ['Fra brev til køb', 'Antal', '% af forrige', 'Faldt fra', 'Handling', 'Brev, mail eller script', 'Hvornår'])
    ws.row_dimensions[hr].height = 30

    r = hr + 1
    # Breve sendt: ét tal, men to dokumenter kan være relevante. Kun det segmentets eget brev vises.
    breve_dok = [(BREV[seg], f'Segment {seg}')]
    lbl(ws, f'A{r}', 'Breve sendt', bold=True)
    fml(ws, f'B{r}', f'=Breve!F{br}')
    lbl(ws, f'E{r}', BUDSKAB[seg] + ' Send som Word-fil i Resights, så flettefelterne udfyldes.')
    c = ws[f'F{r}']; c.value = breve_dok[0][0]; c.hyperlink = LINKS[breve_dok[0][0]]; c.font = LINK_FONT
    lbl(ws, f'G{r}', breve_dok[0][1])
    ws.row_dimensions[r].height = 40
    for col in 'ABCDEFG':
        ws[f'{col}{r}'].alignment = Alignment(wrap_text=True, vertical='center', horizontal='right' if col in 'BCD' else None)
    brev_r = r
    r += 1
    lead_r = r
    forrige = brev_r
    for navn, nogle, handling, dok in TRIN:
        n = max(1, len(dok))
        lbl(ws, f'A{r}', navn, bold=True)
        hard(ws, f'B{r}', tal[nogle])
        if navn == 'Lead':
            fml(ws, f'C{r}', f'=IF(B{brev_r}>0,B{r}/B{brev_r},"–")', fmt=PCT)
            fml(ws, f'D{r}', f'=IF(B{brev_r}>=B{r},B{brev_r}-B{r},"–")')
        else:
            fml(ws, f'C{r}', f'=IF(B{forrige}>0,B{r}/B{forrige},"–")', fmt=PCT)
            fml(ws, f'D{r}', f'=B{forrige}-B{r}')
        lbl(ws, f'E{r}', handling)
        for j, (d, hv) in enumerate(dok):
            c = ws[f'F{r+j}']; c.value = d; c.hyperlink = LINKS[d]; c.font = LINK_FONT
            lbl(ws, f'G{r+j}', hv)
        for rr in range(r, r + n):
            for col in 'ABCDEFG':
                ws[f'{col}{rr}'].alignment = Alignment(wrap_text=True, vertical='center' if n == 1 else 'top', horizontal='right' if col in 'BCD' else None)
            ws.row_dimensions[rr].height = 22 if len(handling) < 60 else 34
        if n > 1:
            for col in 'ABCDE':
                al = ws[f'{col}{r}'].alignment
                ws.merge_cells(f'{col}{r}:{col}{r+n-1}')
                ws[f'{col}{r}'].alignment = Alignment(wrap_text=True, vertical='center', horizontal=al.horizontal)
        forrige = r
        r += n
    if seg == 'B':
        note(ws, f'A{r+1}', 'Segment B er ikke sendt endnu. Leads her kommer fra ejere, der er nået ind på anden vis, for eksempel via julibrevet 2025.')
    if tal['uden_liste']:
        note(ws, f'A{r+1}', f"Heraf {tal['uden_liste']} leads uden for brevlisten (fx Nansensgade 14, hvor der er sendt brev, men adresserne ikke står på listen). De er talt med i Lead og videre, men ikke i Breve sendt.")
        r += 1
    note(ws, f'A{r+2 if seg == "B" else r+1}', f'Hentet fra CRM {ctx["hentet"]:%d-%m-%Y}. Et lead tæller med på alle de trin, det har passeret. Hvert link åbner præcis det brev, den mail eller det script, vi bruger.')
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--ejere', required=True)
    ap.add_argument('--kontakt', required=True)
    ap.add_argument('--mails', required=True)
    ap.add_argument('-o', '--ud', default=str(VAULT / 'Funnel privat boligkøb – A og B.xlsx'))
    a = ap.parse_args()

    def ekstra(wb, ctx):
        funnel_fane(wb, 'A', ctx)
        funnel_fane(wb, 'B', ctx)
        byg_lister(wb, a.ejere, a.kontakt, a.mails)
        del wb['Funnel']
        wb._sheets = [wb[n] for n in ['Funnel A', 'Funnel B', 'Ejerforeninger', 'Ejere', 'Breve']]

    byg(Path(a.ud), ekstra)
    print('skrevet', a.ud)


if __name__ == '__main__':
    main()
