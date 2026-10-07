#!/usr/bin/env python3
"""
Funnel-regnearket i sin enkleste form: to faner.

  1  Funnel  fra ejerforeninger til køb, med handling og link til brev og mails
  2  Breve   hvem der har fået brev, hvilket budskab, og hvor mange gange

Segment A = ejer bor på adressen, segment B = ejer bor et andet sted.
(I brevlisten hedder de A+C og B+D. C og D er de reklamebeskyttede og tælles her med.)

Tallene hentes live fra CRM (/api/admin/funnel-data), brevlisten og BBR-opgørelsen og
skrives som blå tal. Procenter, frafald og bjælker er formler. Ingen adresser, navne
eller mails.

Kør:  python3 scripts/funnel/bygfunnel_simpel.py [versionsnummer]
"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).parent))
from bygfunnel import ROLLUP, VAULT, hent_crm, laes_brevliste, til_dato  # noqa: E402
from mstyle import PCT, _font, fml, hard, header_row, inp, lbl, note, put, title, widths  # noqa: E402

SENDT_A = datetime(2026, 9, 11)
LINKS = json.load(open(Path(__file__).with_name('drive-links.json')))
LINK_FONT = Font(name='Arial', size=10, color='0000FF', underline='single')
BREV = {'A': 'Beboende ejer brev 2', 'B': 'Udlejer brev 1'}
BUDSKAB = {
    'A': 'Kontant bud uden mægler og fremvisninger. Du kan blive boende som lejer.',
    'B': 'Kontant bud, med eller uden lejer, uden mægler og fremvisninger.',
}
HVEM = {'A': 'Ejer bor på adressen', 'B': 'Ejer bor et andet sted'}
SEG = lambda s: 'A' if s in 'AC' else 'B'   # brevlistens A+C → A, B+D → B


def tael(crm, breve):
    """Tæl leads pr. segment og trin. Et lead hører til et brev, hvis adressen matcher."""
    efter_hash = {b['hash']: b for b in breve}
    efter_gade = {}
    for b in breve:
        efter_gade.setdefault(b['gade'], []).append(b)
    trin = ['lead', 'booking', 'aftalt', 'afholdt', 'bud', 'koebt']
    tal = {s: {t: 0 for t in trin} for s in 'AB'}
    for l in crm['leads']:
        if l['test'] or not l['oprettet']:
            continue
        b = efter_hash.get(l['nogle']) if l.get('nogle') else None
        if not b and l.get('nogleGade') and len(efter_gade.get(l['nogleGade'], [])) == 1:
            b = efter_gade[l['nogleGade']][0]
        if not b or SEG(b['segment']) == 'B' or til_dato(l['oprettet']) < SENDT_A:
            continue   # B er ikke sendt, og leads før brevet hører ikke til det
        d = lambda k: bool(l[k])
        koebt, ikke_enige = d('tKoebt'), d('tIkkeEnige')
        bud = d('tBud') or ikke_enige or koebt
        afholdt = d('tAfholdt') or bud
        aftalt = d('tAftalt') or afholdt
        for k, v in (('lead', True), ('booking', d('tBooking')), ('aftalt', aftalt), ('afholdt', afholdt), ('bud', bud), ('koebt', koebt)):
            tal['A'][k] += bool(v)
    return tal


def byg(ud):
    crm, breve = hent_crm(), laes_brevliste()
    rollup = {f['forening']: f for f in json.load(open(ROLLUP))['foreninger']}
    tal = tael(crm, breve)
    hentet = til_dato(crm['genereret'])
    antal = {s: sum(1 for b in breve if SEG(b['segment']) == s) for s in 'AB'}
    besk = {s: sum(1 for b in breve if b['segment'] == c) for s, c in (('A', 'C'), ('B', 'D'))}
    fik25 = {s: sum(1 for b in breve if SEG(b['segment']) == s and b['jul'] == 'Ja') for s in 'AB'}
    # Top: foreninger
    foren = crm['foreninger']
    udvalgt = [f for f in foren if f['status'] == 'maalgruppe']
    stoerrelse = 0
    for f in udvalgt:
        ru = rollup.get(f['gade'])
        if ru and ru['foreningNavn'] == f['navn']:
            stoerrelse += ru['iMaalgruppe']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Breve ───────────────────────────────────────────────────────────
    wb_ = wb.create_sheet('Breve')
    title(wb_, 'Breve', 'Hvem der har fået brev, hvilket budskab, og hvor mange gange. Skriv sendedatoen i de gule felter, når et segment sendes.', 13)
    widths(wb_, {'A': 9, 'B': 24, 'C': 9, 'D': 12, 'E': 12, 'F': 9, 'G': 12, 'H': 9, 'I': 9, 'J': 9, 'K': 8, 'L': 46, 'M': 22})
    header_row(wb_, 4, ['Segment', 'Hvem', 'Antal', 'Heraf reklame-beskyttet', 'Sendt den', 'Sendt', 'Fik brev jul 2025', '0 breve', '1 brev', '2 breve', 'Leads', 'Budskab', 'Brev'])
    wb_.row_dimensions[4].height = 42
    for j, s in enumerate('AB'):
        r = 5 + j
        lbl(wb_, f'A{r}', s, bold=True)
        lbl(wb_, f'B{r}', HVEM[s])
        hard(wb_, f'C{r}', antal[s])
        hard(wb_, f'D{r}', besk[s])
        inp(wb_, f'E{r}', SENDT_A if s == 'A' else None, fmt='dd-mm-yyyy')
        fml(wb_, f'F{r}', f'=IF(E{r}="",0,C{r})')
        hard(wb_, f'G{r}', fik25[s])
        fml(wb_, f'H{r}', f'=C{r}-I{r}-J{r}')
        fml(wb_, f'I{r}', f'=IF(E{r}="",G{r},C{r}-G{r})')
        fml(wb_, f'J{r}', f'=IF(E{r}="",0,G{r})')
        hard(wb_, f'K{r}', tal['A']['lead'] if s == 'A' else 0)
        lbl(wb_, f'L{r}', BUDSKAB[s])
        wb_[f'L{r}'].alignment = Alignment(wrap_text=True, vertical='center')
        c = wb_[f'M{r}']; c.value = BREV[s]; c.hyperlink = LINKS[BREV[s]]; c.font = LINK_FONT
        wb_.row_dimensions[r].height = 32
    lbl(wb_, 'A7', 'I alt', bold=True)
    for c in 'CDFGHIJK':
        fml(wb_, f'{c}7', f'=SUM({c}5:{c}6)', bold=True)
    note(wb_, 'A9', 'Segment A er ejer bor på adressen. Segment B er ejer bor et andet sted. Reklamebeskyttede er med i begge.')
    note(wb_, 'A10', 'Brevene til B lander hos lejeren, ikke ejeren, fordi ejerens postadresse ikke er i datasættet.')
    note(wb_, 'A11', f'Antal breve pr. modtager er brevet i juli 2025 (hvis modtageren var med) plus brevet i 2026 (hvis sendt). Leads hentet fra CRM {hentet:%d-%m-%Y}.')

    # ── Funnel ──────────────────────────────────────────────────────────
    ws = wb.create_sheet('Funnel', 0)
    title(ws, 'Funnel: fra ejerforeninger til køb', f'Hentet fra CRM {hentet:%d-%m-%Y}. Et lead tæller med på alle de trin, det har passeret. Hvert link åbner præcis det brev, den mail eller det script, vi bruger.', 9)
    widths(ws, {'A': 34, 'B': 11, 'C': 38, 'D': 12, 'E': 10, 'F': 11, 'G': 46, 'H': 32, 'I': 24})

    # Top: ejerforeninger → segmenter
    header_row(ws, 4, ['Fra ejerforeninger til brev', 'Antal', 'Tragt', '% af forrige', '', '', 'Hvad tallet er'])
    ws.row_dimensions[4].height = 26
    top = [
        ('Ejerforeninger identificeret', len(foren), 'foren', 'Alle foreninger i registret på /foreninger'),
        ('heraf udvalgt', len(udvalgt), 'foren', 'Status Målgruppe: dem vi vil købe i'),
        ('heraf størrelse 20–80 kvm', stoerrelse, 'bolig', 'Boliger i de udvalgte, målt i BBR'),
        ('kan få brev', f'=Breve!C7', 'bolig', 'Ejeren er kendt i Resights, og koncernens egne lejligheder er trukket fra'),
        ('Segment A: ejer bor der', '=Breve!C5', 'bolig', 'Brevet når ejeren selv'),
        ('Segment B: ejer bor ikke der', '=Breve!C6', 'bolig', 'Brevet lander hos lejeren'),
    ]
    for i, (navn, v, enhed, kom) in enumerate(top):
        r = 5 + i
        indent = 1 if i in (1, 2) else 2 if i >= 4 else 0
        lbl(ws, f'A{r}', navn, bold=i in (0, 3), indent=indent)
        (fml if isinstance(v, str) else hard)(ws, f'B{r}', v)
        roed = '$B$5' if enhed == 'foren' else '$B$7'   # bjælkerne skaleres inden for samme enhed
        fml(ws, f'C{r}', f'=IF({roed}>0,REPT("█",MAX(IF(B{r}>0,1,0),ROUND(28*B{r}/{roed},0))),"")', fmt='@')
        ws[f'C{r}'].font = _font(False, '1E293B', 10)
        if i in (1, 3):
            fml(ws, f'D{r}', f'=IF(B{r-1}>0,B{r}/B{r-1},"–")', fmt=PCT)
        elif i in (4, 5):
            fml(ws, f'D{r}', f'=IF($B$8>0,B{r}/$B$8,"–")', fmt=PCT)
        note(ws, f'G{r}', kom)
        for c in 'BCD':
            ws[f'{c}{r}'].alignment = Alignment(horizontal=('center' if c == 'C' else 'right'), vertical='center')
        ws.row_dimensions[r].height = 22

    # Tragten fra brev til køb
    hr = 13
    header_row(ws, hr, ['Fra brev til køb', 'A: ejer bor', 'Tragt (A)', '% af forrige', 'Faldt fra', 'B: ejer bor ikke', 'Handling', 'Brev, mail eller script', 'Hvornår'])
    ws.row_dimensions[hr].height = 30
    A_ = tal['A']
    rows = [
        ('Breve sendt', '=Breve!F5', '=Breve!F6', 'A: ' + BUDSKAB['A'] + ' B: ' + BUDSKAB['B'] + ' Send som Word-fil i Resights, så flettefelterne udfyldes.',
         [('Beboende ejer brev 2', 'Segment A'), ('Udlejer brev 1', 'Segment B')]),
        ('Lead', A_['lead'], 0, 'Send booking-mailen senest dagen efter, at de har brugt beregneren.',
         [('Mail 01 Booking', 'Senest dagen efter')]),
        ('Booking sendt', A_['booking'], 0, 'Følg op i samme tråd. Uden svar efter tredje opfølgning parkeres leadet.',
         [('Mail 02 Opfølgning 1 dag 3', 'Dag 3'), ('Mail 03 Opfølgning 2 dag 8', 'Dag 8'), ('Mail 04 Opfølgning 3 dag 15', 'Dag 15')]),
        ('Besigtigelse aftalt', A_['aftalt'], 0, 'Bekræft tiden skriftligt, og påmind dagen før. Kør derud med scriptet.',
         [('Mail 05 Bekræftelse', 'Straks'), ('Mail 06 Påmindelse', 'Dagen før'), ('Script Besigtigelse', 'På dagen')]),
        ('Besigtigelse afholdt', A_['afholdt'], 0, 'Giv et skriftligt bud senest dagen efter.',
         [('Mail 07 Tak for i dag', 'Samme dag')]),
        ('Bud afgivet', A_['bud'], 0, 'Send kontantbud, som gælder i 30 dage, og følg op. Tæller også dem, der siden blev uenige om prisen.',
         [('Mail 08 Bud', 'Dagen efter besigtigelsen'), ('Mail 09 Bud-opfølgning dag 3', 'Dag 3')]),
        ('Købt', A_['koebt'], 0, 'Send de næste skridt, når I er blevet enige.',
         [('Mail 14 Sådan foregår handlen', 'Når I er enige')]),
    ]
    first = hr + 1
    lead_r = first + len(rows[0][4])   # leadrækken kommer efter brevenes linkrækker
    r = first
    for i, (navn, a, b, handling, dok) in enumerate(rows):
        antal_r = max(1, len(dok))
        lbl(ws, f'A{r}', navn, bold=True)
        (fml if isinstance(a, str) else hard)(ws, f'B{r}', a)
        (fml if isinstance(b, str) else hard)(ws, f'F{r}', b)
        if i == 0:
            fml(ws, f'C{r}', f'=IF(B{r}>0,"Respons: "&ROUND(100*B{lead_r}/B{r},1)&" % blev til et lead","")', fmt='@')
            ws[f'C{r}'].font = _font(False, '64748B', 9, True)
        else:
            fml(ws, f'C{r}', f'=IF($B${lead_r}>0,REPT("█",MAX(IF(B{r}>0,1,0),ROUND(28*B{r}/$B${lead_r},0))),"")', fmt='@')
            ws[f'C{r}'].font = _font(False, '1E293B', 10)
            fml(ws, f'D{r}', f'=IF(B{r-antal_prev}>0,B{r}/B{r-antal_prev},"–")', fmt=PCT)
            fml(ws, f'E{r}', f'=B{r-antal_prev}-B{r}')
        lbl(ws, f'G{r}', handling)
        for j, (d, hvornaar) in enumerate(dok):
            rr = r + j
            c = ws[f'H{rr}']; c.value = d; c.hyperlink = LINKS[d]; c.font = LINK_FONT
            lbl(ws, f'I{rr}', hvornaar)
        for rr in range(r, r + antal_r):
            for c in 'ABCDEFGHI':
                ind = ws[f'{c}{rr}'].alignment.indent
                ws[f'{c}{rr}'].alignment = Alignment(
                    wrap_text=True, vertical='center' if antal_r == 1 else 'top',
                    horizontal=('center' if c == 'C' else 'right' if c in 'BDEF' else None), indent=ind)
            ws.row_dimensions[rr].height = 22
        if antal_r > 1:
            for c in 'ABCDEFG':
                al = ws[f'{c}{r}'].alignment
                ws.merge_cells(f'{c}{r}:{c}{r+antal_r-1}')
                ws[f'{c}{r}'].alignment = Alignment(wrap_text=True, vertical='center', horizontal=al.horizontal, indent=al.indent)
        if len(handling) > 60:
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height, 34)
        if i == 0:
            ws.row_dimensions[r].height = 40
            ws.row_dimensions[r + 1].height = 40
        antal_prev = antal_r
        r += antal_r
    note(ws, f'A{r+1}', 'Bjælkerne viser segment A i forhold til antal leads. Segment B er ikke sendt endnu. Foreninger og boliger har hver sin skala.')
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    wb.save(ud)


if __name__ == '__main__':
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 3
    ud = VAULT / f'Funnel boligberegner – v{n}.xlsx'
    byg(ud)
    print('skrevet', ud)
