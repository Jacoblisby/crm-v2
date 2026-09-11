"""
Brevskabeloner som Word-filer til Resights' brevmodul.

  python3 scripts/brev/brev_docx.py            # alle breve
  python3 scripts/brev/brev_docx.py udlejer    # kun ét

Tre ting, lært af første runde (11.09.2026):

· Resights fletter kun i Word-filer. En PDF bliver trykt, som den er — med
  {{fornavn}} stående i klar tekst.
· Resights udfylder skabelonens EGEN adresseblok ({{att_navn}},
  {{vej_og_nummer}}, {{postnummer}} {{by}}). Den skal med.
· Hvert flettefelt skrives som ét samlet tekststykke. Et felt, der er delt
  op i Word-filen («{{forn» + «avn}}»), kan en flettemotor ikke finde.
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VAULT = Path.home() / 'Desktop/Claude Vault/Projects/Brevkampagne ejerforeninger'
QR = VAULT / 'QR saelg.365ejendom.dk.png'
PETROL = RGBColor(0x14, 0x5D, 0x5F)
TEKST = RGBColor(0x1F, 0x2A, 0x2B)
GRAA = RGBColor(0x5B, 0x68, 0x68)
HVID = RGBColor(0xFF, 0xFF, 0xFF)

SAMMEN = ('Siden 2020 har vi købt mere end 85 boliger direkte af sælgere på Sjælland, og vi køber for '
          'at eje og udleje, ikke for at sælge videre.')

BREVE = {
    # Segment A + C: ejeren bor selv i lejligheden. Sendt 11.09.2026.
    'beboende': {
        'fil': 'Beboende ejer brev 2.docx',
        'overskrift': 'Et kontant bud på din lejlighed',
        'under': 'Uden mægler. Uden fremvisninger. Og du kan blive boende.',
        'indledning': [
            [('Hej ', 0), ('{{fornavn}}', 0)],
            [('Vi vil gerne købe din lejlighed på ', 0), ('{{ejendomsadresse}}', 0),
             ('. Kontant og direkte, uden mægler.', 0)],
            [('Du får dette brev, fordi vi allerede ejer lejligheder i din ejerforening. Vi kender '
              'bygningen og foreningen, så vi kan give dig et konkret bud hurtigt. ' + SAMMEN, 0)],
        ],
        'kolonne': 'Almindeligt salg',
        'tabel': [
            ('Salær og udbud', 'Typisk 50–80.000 kr.', '0 kr.'),
            ('Fremvisninger', 'Åbent hus og fremmede i dit hjem', 'Ingen. Lejligheden kommer ikke til salg'),
            ('Betaling', 'Afhænger af købers bank', 'Kontant. Intet bankforbehold'),
            ('Overtagelse', 'Når køberen er klar', 'Når du er klar'),
            ('Bagefter', 'Du skal flytte', 'Du kan ofte blive boende som lejer'),
        ],
        'situation': ('Måske har du fundet din drømmebolig og vil hurtigt videre, uden at vente på en køber. '
                      'Måske vil du have friværdien ud, men blive i dit hjem som lejer. Eller måske er det noget '
                      'helt tredje. Du bestemmer selv tidspunktet, i dag eller om et halvt år, og vi gennemgår '
                      'alle vilkår med dig, før du beslutter noget.'),
        'boks_titel': 'Se dit bud på et par minutter',
        'boks': [
            [('Scan koden, eller gå ind på ', 0), ('saelg.365ejendom.dk', 1),
             (' og skriv din adresse. Så får du et foreløbigt bud med det samme. Det er gratis og '
              'forpligter dig ikke til noget.', 0)],
            [('Vil du hellere tale med et menneske, så ring direkte til mig på ', 0), ('61 78 90 71', 1), ('.', 0)],
        ],
        'ps': ('Du behøver ikke have planer om at sælge for at scanne koden. Det kan være rart bare at vide, '
               'hvad din lejlighed er værd i dag. Buddet er dit at tage stilling til, nu, senere eller aldrig.'),
    },

    # Segment B + D: ejeren bor et andet sted — udlejere, selskaber, forældrekøb.
    'udlejer': {
        'fil': 'Udlejer brev 1.docx',
        'overskrift': 'Et kontant bud på din lejlighed',
        'under': 'Med eller uden lejer. Uden mægler. Uden fremvisninger.',
        'indledning': [
            [('Hej ', 0), ('{{fornavn}}', 0)],
            [('Vi vil gerne købe din lejlighed på ', 0), ('{{ejendomsadresse}}', 0),
             ('. Kontant og direkte, med eller uden lejer.', 0)],
            [('Du får dette brev, fordi vi allerede ejer og udlejer lejligheder i samme ejerforening. Vi kender '
              'bygningen, foreningen og lejeniveauet, så vi kan give dig et konkret bud hurtigt. ' + SAMMEN, 0)],
        ],
        'kolonne': 'Salg gennem mægler',
        'tabel': [
            ('Lejeren', 'Mange købere fravælger en udlejet bolig', 'Bliver boende. Vi overtager lejekontrakten'),
            ('Fremvisninger', 'Kræver lejerens samarbejde', 'Ingen'),
            ('Salær og udbud', 'Typisk 50–80.000 kr.', '0 kr.'),
            ('Betaling', 'Afhænger af købers bank', 'Kontant. Intet bankforbehold'),
            ('Overtagelse', 'Når køberen er klar', 'Når du er klar'),
        ],
        'situation': ('Måske har du købt lejligheden til dit barn, som nu flytter videre. Måske er du træt af at '
                      'være udlejer, eller vil hellere bruge pengene på noget andet. Eller måske er det noget helt '
                      'tredje. Vi køber med eller uden lejer, og du bestemmer selv tidspunktet. Ejer du flere '
                      'lejligheder, giver vi gerne et samlet bud på dem alle.'),
        'boks_titel': 'Få et bud',
        'boks': [
            [('Ring direkte til mig på ', 0), ('61 78 90 71', 1), (', eller skriv til ', 0),
             ('administration@365ejendom.dk', 1), ('. Så regner vi på lejligheden, som den er i dag, med lejer '
              'og lejekontrakt.', 0)],
            [('Du kan også scanne koden eller gå ind på ', 0), ('saelg.365ejendom.dk', 1),
             (' for et foreløbigt bud med det samme. Det er gratis og forpligter dig ikke til noget.', 0)],
        ],
        'ps': ('Du behøver ikke have planer om at sælge for at høre, hvad vi vil give. Det kan være rart bare at '
               'vide, hvad lejligheden er værd i dag, også med lejer.'),
    },
}


# ── Hjælpere ────────────────────────────────────────────────────────────
def afsnit(doc, dele, efter=6, foer=0, size=None, color=None, container=None, linje=None):
    """dele: str eller liste af (tekst, fed). Hver del bliver ét run — så flettefelter holdes samlet."""
    p = (container or doc).add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(efter), Pt(foer)
    if linje:
        pf.line_spacing = linje
    for tekst, fed in ([(dele, 0)] if isinstance(dele, str) else dele):
        r = p.add_run(tekst)
        r.bold = bool(fed)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p


def celle_fyld(celle, hex_):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_)
    celle._tc.get_or_add_tcPr().append(shd)


def rammer(tabel, hex_, sz=4):
    b = OxmlElement('w:tblBorders')
    for kant in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{kant}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), hex_)
        b.append(e)
    tabel._tbl.tblPr.append(b)


def celle_margin(tabel, top=50, bund=50, venstre=100, hoejre=100):
    m = OxmlElement('w:tblCellMar')
    for kant, v in (('top', top), ('bottom', bund), ('left', venstre), ('right', hoejre)):
        e = OxmlElement(f'w:{kant}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tabel._tbl.tblPr.append(m)


def lås_bredder(tabel, bredder):
    """Faste kolonnebredder. Ellers fordeler Word/LibreOffice kolonnerne ligeligt."""
    tabel.autofit = False
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed')
    tabel._tbl.tblPr.append(layout)
    for kol, w in zip(tabel.columns, bredder):
        kol.width = w
        for c in kol.cells:
            c.width = w


def celle_tekst(celle, tekst, fed=False, color=None, size=9.5):
    p = celle.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(tekst); r.bold = fed; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


# ── Byg ─────────────────────────────────────────────────────────────────
def byg(b):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(22)
    s.top_margin, s.bottom_margin = Mm(20), Mm(16)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEKST
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.3

    # Adresseblok — Resights udfylder den med modtagerens navn og adresse.
    afsnit(doc, [('{{att_navn}}', 0)], efter=0, linje=1.15)
    afsnit(doc, [('{{vej_og_nummer}}', 0)], efter=0, linje=1.15)
    afsnit(doc, [('{{postnummer}}', 0), (' ', 0), ('{{by}}', 0)], efter=16, linje=1.15)

    afsnit(doc, [(b['overskrift'], 1)], efter=2, size=18, color=PETROL, linje=1.1)
    afsnit(doc, b['under'], efter=11, size=11, color=PETROL)
    for i, dele in enumerate(b['indledning']):
        afsnit(doc, dele, efter=10 if i == len(b['indledning']) - 1 else 6)

    t = doc.add_table(rows=1 + len(b['tabel']), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rammer(t, 'C9D8D8'); celle_margin(t); lås_bredder(t, (Mm(40), Mm(62), Mm(64)))
    h = t.rows[0].cells
    celle_fyld(h[0], 'F2F6F6'); celle_fyld(h[1], 'F2F6F6'); celle_fyld(h[2], '145D5F')
    celle_tekst(h[1], b['kolonne'], fed=True, color=GRAA)
    celle_tekst(h[2], 'Salg til 365 Ejendomme', fed=True, color=HVID)
    for i, (a, x, y) in enumerate(b['tabel'], start=1):
        cs = t.rows[i].cells
        celle_tekst(cs[0], a, fed=True)
        celle_tekst(cs[1], x, color=GRAA)
        celle_fyld(cs[2], 'EAF3F3'); celle_tekst(cs[2], y, fed=True)

    afsnit(doc, [('Hvad end din situation er', 1)], efter=1, foer=11)
    afsnit(doc, b['situation'], efter=11)

    bx = doc.add_table(rows=1, cols=2)
    bx.alignment = WD_TABLE_ALIGNMENT.CENTER
    rammer(bx, '145D5F', sz=12); celle_margin(bx, top=130, bund=130, venstre=170, hoejre=170)
    lås_bredder(bx, (Mm(134), Mm(32)))
    v, hq = bx.rows[0].cells
    p0 = v.paragraphs[0]; p0.paragraph_format.space_after = Pt(3)
    r = p0.add_run(b['boks_titel']); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = PETROL
    for i, dele in enumerate(b['boks']):
        afsnit(doc, dele, efter=0 if i == len(b['boks']) - 1 else 4, container=v)
    hq.vertical_alignment = 1
    hq.paragraphs[0].add_run().add_picture(str(QR), width=Mm(26))

    afsnit(doc, 'Venlig hilsen', efter=4, foer=13)
    afsnit(doc, 'Jacob Fast Lisby', efter=0)
    afsnit(doc, '365ejendom ApS', efter=14)
    afsnit(doc, [('P.S. ', 1), (b['ps'], 0)], efter=0)

    ud = VAULT / b['fil']
    doc.save(ud)
    return ud


if __name__ == '__main__':
    valg = sys.argv[1:] or list(BREVE)
    for navn in valg:
        print('Gemt:', byg(BREVE[navn]))
