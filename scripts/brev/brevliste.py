"""
Brevliste til en brevrunde — flettefil klar til tryk.

Læser ejere.json (Resights), kobler hver lejlighed til sin forening via
bbr-enheder.json + bbr-forening-map.json, og trækker koncernens egne
lejligheder fra ud fra fanen «Vi ejer (Resights)» i brevstatus-arket.

Alle modtagere kommer med — også de reklamebeskyttede. Hvem der faktisk
skal have brev, afgøres ved print ud fra kolonnen «Segment», ikke her.
Det er med vilje: beslutningen er forretningens, og filen skal vise den
tydeligt i stedet for at træffe den i stilhed.

  A  må kontaktes      ejer bor på adressen   brevet når ejeren
  B  må kontaktes      ejer bor et andet sted brevet når lejeren
  C  reklamebeskyttet  ejer bor på adressen
  D  reklamebeskyttet  ejer bor et andet sted

Ejernavn og ejeradresse findes ikke i datasættet (se scripts/bbr/ejere.py),
så alle breve adresseres «Til ejeren» på lejlighedens egen adresse.

Kør:  python3 scripts/brev/brevliste.py
"""
import json, re, sys
from datetime import date
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROD = Path(__file__).resolve().parents[2]
DATA = ROD / 'src/lib/data'
VAULT = Path.home() / 'Desktop/Claude Vault/Projects/Brevkampagne ejerforeninger'
STATUS = VAULT / 'Ejerforeninger breve Status.xlsx'
UD = VAULT / f'Brevliste runde 2 {date.today():%Y-%m}.xlsx'
QR = 'https://saelg.365ejendom.dk'
KVM_MIN, KVM_MAX = 20, 80

SEGMENT = {
    ('Nej', True):  ('A', 'Må kontaktes · ejer bor der'),
    ('Nej', False): ('B', 'Må kontaktes · ejer bor et andet sted'),
    ('Ja', True):   ('C', 'Reklamebeskyttet · ejer bor der'),
    ('Ja', False):  ('D', 'Reklamebeskyttet · ejer bor et andet sted'),
}
FARVE = {'A': 'D9EAD3', 'B': 'FFF2CC', 'C': 'F4CCCC', 'D': 'EAD1DC'}


def norm(s):
    return re.sub(r'[\s.]+', ' ', s.lower()).strip()


def main():
    enh = json.load(open(DATA / 'ejere.json'))['enheder']
    bbr = json.load(open(DATA / 'bbr-enheder.json'))['enheder']
    kob = json.load(open(DATA / 'bbr-forening-map.json'))['kobling']
    adr2ejd = {(norm(r['adresse']), str(r['postnr'])): r['ejendomBfe'] for r in bbr}

    wb_s = openpyxl.load_workbook(STATUS, read_only=True, data_only=True)
    vi_ejer = {int(r[0]) for r in list(wb_s['Vi ejer (Resights)'].iter_rows(values_only=True))[1:]
               if isinstance(r[0], (int, float))}
    sendt = {int(c) for r in wb_s['Sendt 24. jul9 2025'].iter_rows(values_only=True)
             for c in r if isinstance(c, (int, float)) and c > 100000}

    maal = [u for u in enh if u.get('erBolig') and u.get('kvm') and KVM_MIN <= u['kvm'] <= KVM_MAX]

    rækker, egne = [], []
    for u in maal:
        dele = u['adresse'].split(', ')
        ejd = adr2ejd.get((norm(', '.join(dele[:2])), str(u['postnr'])))
        forening = kob.get(str(ejd))
        if not forening:
            sys.exit(f'Ingen forening for {u["adresse"]} — ret bbr-forening-map.json')
        if u['bfe'] in vi_ejer:
            egne.append((u, forening))
            continue
        rb = 'Ja' if u['reklamebeskyttet'] != 'Nej' else 'Nej'   # Ukendt behandles som Ja
        seg, segtekst = SEGMENT[(rb, bool(u['ejerBorDer']))]
        rækker.append({
            'Segment': seg,
            'Segment (tekst)': segtekst,
            'Brevet når': 'Ejer' if u['ejerBorDer'] else 'Lejer / beboer',
            'Reklamebeskyttet': u['reklamebeskyttet'],
            'Ejertype': u['ejerType'],
            'Forening': forening,
            'Modtager': 'Til ejeren',
            'Adresselinje 1': ', '.join(dele[:2]),
            'Adresselinje 2': dele[2] if len(dele) == 4 else '',
            'Postnr og by': dele[-1],
            'Kvm': u['kvm'],
            'Fik brev jul 2025': 'Ja' if u['bfe'] in sendt else 'Nej',
            'BFE': u['bfe'],
            'QR-link': QR,
        })
    rækker.sort(key=lambda r: (r['Segment'], r['Forening'], r['Adresselinje 1']))

    wb = openpyxl.Workbook()
    fed = Font(bold=True)

    # ── Overblik ──
    ov = wb.active
    ov.title = 'Overblik'
    ov.append([f'Brevliste runde 2 · genereret {date.today():%d.%m.%Y}']); ov['A1'].font = Font(bold=True, size=14)
    ov.append([])
    ov.append(['Målgruppe (bolig, 20–80 kvm)', len(maal)])
    ov.append(['− koncernens egne lejligheder', len(egne)])
    ov.append(['= mulige modtagere', len(rækker)])
    ov.append([])
    ov.append(['Segment', 'Beskrivelse', 'Antal', 'Heraf fik brev jul 2025']); [setattr(c, 'font', fed) for c in ov[ov.max_row]]
    for (rb, bor), (seg, tekst) in sorted(SEGMENT.items(), key=lambda x: x[1][0]):
        sub = [r for r in rækker if r['Segment'] == seg]
        ov.append([seg, tekst, len(sub), sum(r['Fik brev jul 2025'] == 'Ja' for r in sub)])
        ov.cell(ov.max_row, 1).fill = PatternFill('solid', fgColor=FARVE[seg])
    ov.append([])
    ov.append(['Pr. forening', 'A', 'B', 'C', 'D', 'I alt']); [setattr(c, 'font', fed) for c in ov[ov.max_row]]
    for f in sorted({r['Forening'] for r in rækker}):
        n = {s: sum(1 for r in rækker if r['Forening'] == f and r['Segment'] == s) for s in 'ABCD'}
        ov.append([f, n['A'], n['B'], n['C'], n['D'], sum(n.values())])
    ov.append([])
    for linje in [
        'Noter',
        '· C og D er reklamebeskyttede (CPR for privatpersoner, CVR for selskaber). «Ukendt» (5 stk.) tælles som beskyttet.',
        '· B og D: ejeren bor et andet sted. Brevet lander hos lejeren/beboeren — ejerens postadresse er ikke i datasættet.',
        '· Alle adresseres «Til ejeren» på lejlighedens adresse. Ejernavn gemmes bevidst ikke.',
        '· QR peger på forsiden (saelg.365ejendom.dk). Personlige koder (/k/<kode>) er ikke bygget endnu.',
        '· Koncernens egne lejligheder er trukket fra ud fra fanen «Vi ejer (Resights)» — se fanen «Fratrukket (egne)».',
    ]:
        ov.append([linje])
    ov.cell(ov.max_row - 5, 1).font = fed
    ov.column_dimensions['A'].width = 34; ov.column_dimensions['B'].width = 42
    for c in 'CDEF': ov.column_dimensions[c].width = 12

    # ── Flettefil ──
    ff = wb.create_sheet('Flettefil')
    kol = list(rækker[0])
    ff.append(kol); [setattr(c, 'font', fed) for c in ff[1]]
    for r in rækker:
        ff.append([r[k] for k in kol])
        ff.cell(ff.max_row, 1).fill = PatternFill('solid', fgColor=FARVE[r['Segment']])
    ff.freeze_panes = 'B2'
    ff.auto_filter.ref = f'A1:{get_column_letter(len(kol))}{ff.max_row}'
    for i, k in enumerate(kol, 1):
        ff.column_dimensions[get_column_letter(i)].width = max(len(k), *(len(str(r[k])) for r in rækker[:300])) + 2

    # ── Fratrukket ──
    fr = wb.create_sheet('Fratrukket (egne)')
    fr.append(['BFE', 'Adresse', 'Forening', 'Ejertype i Resights nu']); [setattr(c, 'font', fed) for c in fr[1]]
    for u, f in sorted(egne, key=lambda x: x[0]['adresse']):
        fr.append([u['bfe'], u['adresse'], f, u['ejerType']])
    for c, w in zip('ABCD', (10, 48, 24, 22)): fr.column_dimensions[c].width = w

    wb.save(UD)
    print(f'Gemt: {UD}')
    print(f'målgruppe {len(maal)}  − egne {len(egne)}  = {len(rækker)}')
    for s in 'ABCD':
        print(f'  {s}: {sum(1 for r in rækker if r["Segment"] == s)}')
    privat_egne = [u for u, _ in egne if u['ejerType'] != 'Selskab']
    if privat_egne:
        print(f'OBS: {len(privat_egne)} af «egne» står nu med ejertype ≠ Selskab i Resights — måske solgt siden «Vi ejer»-listen blev lavet.')


if __name__ == '__main__':
    main()
